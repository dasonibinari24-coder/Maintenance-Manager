"""Fill only the original, supplied HWPX forms at their original table cells."""
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from io import BytesIO
import json
import re
from datetime import datetime
from zoneinfo import ZoneInfo
from urllib.parse import parse_qs, urlparse
import zipfile
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw, ImageFont
import fitz
import olefile
import openpyxl
from qa_backend import QAConfigurationError, ask as ask_qa

ROOT = Path(__file__).parent
NS = {'hp': 'http://www.hancom.co.kr/hwpml/2011/paragraph'}
ET.register_namespace('hp', NS['hp'])


def template_for(kind):
    if kind == 'form10':
        return next(p for p in ROOT.glob('*.hwpx') if p.name.endswith('선임ㆍ해임 신고서.hwpx'))
    return next(p for p in ROOT.glob('*.hwpx') if '선임신고증명서 발급신청서' in p.name)


def admin_template_for(kind):
    """Use only the two HWPX originals supplied for the administrator workflow."""
    if kind == 'draft':
        return next(p for p in ROOT.glob('*.hwpx') if '수리 알림(TEST)' in p.name)
    return template_for('form12')


def fill_admin_draft(data):
    """Populate only the blanks and placeholders in the supplied 수리 알림 HWPX."""
    replacements = {
        '여울시티 2차 구분소유자 대표 귀중': f"{data.get('ownerName', '')} 대표 {data.get('ownerRepresentative', '')} 귀중".strip(),
        '정보통신설비 유지보수 관리자 선임 신고서 수리 알림 [주소_건축물명]':
            f"정보통신설비 유지보수 관리자 선임 신고서 수리 알림 [{data.get('buildingAddress', '')}_{data.get('buildingName', '')}]",
        '- 상호(명칭) :': f"- 상호(명칭) :\u00a0{data.get('ownerName', '')}",
        '- 대표자 :': f"- 대표자 :\u00a0{data.get('ownerRepresentative', '')}",
        '- 연면적 :': f"- 연면적 :\u00a0{data.get('buildingArea', '')}㎡",
        '- 용도 :': f"- 용도 :\u00a0{data.get('buildingUse', '')}",
        '- 주소 :': f"- 주소 :\u00a0{data.get('buildingAddress', '')}",
        '선임 / 000 / 0급 /0000-00-00':
            f"선임 / {data.get('managerName', '')} / {data.get('managerGrade', '')} / {data.get('appointmentDate', '')}",
    }
    buffer = BytesIO()
    with zipfile.ZipFile(admin_template_for('draft'), 'r') as source, zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as destination:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == 'Contents/section0.xml':
                root = ET.fromstring(content)
                for text in root.findall('.//hp:t', NS):
                    value = text.text or ''
                    key = value.strip()
                    if key in replacements:
                        text.text = replacements[key]
                    elif '민원접수번호: 00000, 2026. 0. 00.' in value:
                        # 결재 전 수리 알림에는 시행일을 넣지 않는다.
                        text.text = value.replace(', 2026. 0. 00.', '')
                content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            destination.writestr(info, content)
    return buffer.getvalue()


def admin_document(kind, data=None):
    """Return editable HWPX, never a reconstructed PDF.

    The draft remains in its supplied pre-approval state.  The certificate
    application is the supplied HWPX form with only its 작성일 set to today.
    """
    data = data or {}
    if kind == 'draft':
        return fill_admin_draft(data)
    today = datetime.now(ZoneInfo('Asia/Seoul')).date().isoformat()
    return fill_form12({**data, 'reportDate': today})


def admin_preview(kind, data=None):
    """Show the supplied HWPX preview for the same document being downloaded."""
    data = data or {}
    if kind == 'certificate':
        today = datetime.now(ZoneInfo('Asia/Seoul')).date().isoformat()
        return generated_preview('form12', {**data, 'reportDate': today})
    with zipfile.ZipFile(admin_template_for('draft')) as source:
        image = Image.open(BytesIO(source.read('Preview/PrvImage.png'))).convert('RGB')
    draw = ImageDraw.Draw(image)
    # These are the blank/data lines in the supplied 수리 알림 preview image.
    # Erase only its example values; headings, approval boxes and legal text stay
    # exactly as provided in the source HWPX.
    fields = (
        ((110, 179, 650, 204), f"{data.get('ownerName', '')} 대표 {data.get('ownerRepresentative', '')} 귀중"),
        ((110, 226, 680, 250), f"정보통신설비 유지보수 관리자 선임 신고서 수리 알림 [{data.get('buildingAddress', '')}_{data.get('buildingName', '')}]"),
        ((205, 399, 675, 421), data.get('ownerName', '')),
        ((185, 423, 675, 445), data.get('ownerRepresentative', '')),
        ((185, 493, 675, 515), f"{data.get('buildingArea', '')}㎡"),
        ((165, 516, 675, 538), data.get('buildingUse', '')),
        ((165, 539, 675, 562), data.get('buildingAddress', '')),
        ((112, 608, 675, 632), f"선임 / {data.get('managerName', '')} / {data.get('managerGrade', '')} / {data.get('appointmentDate', '')}"),
    )
    for box, value in fields:
        draw.rectangle(box, fill=(255, 255, 255))
        draw_value(draw, (box[0], box[1]), value, box[2] - box[0], size=14)
    output = BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def duplicate_registry(query=''):
    """Read the explicitly supplied, unlocked test registry before archived files."""
    candidates = sorted(ROOT.glob('*test*.xlsx')) or sorted(ROOT.glob('*중복선임 여부 대상자 조회 명단*.xlsx'))
    if not candidates:
        return {'files': [], 'locked': False, 'fields': [], 'rows': [], 'message': '조회명단 파일이 없습니다.'}
    path = candidates[0]
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.active
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        header_index = next((index for index, row in enumerate(values) if any(str(value or '').strip() == '성명' for value in row)), None)
        if header_index is None:
            raise ValueError('열 제목(성명)을 찾을 수 없습니다.')
        columns = [(index, str(value).strip()) for index, value in enumerate(values[header_index]) if str(value or '').strip()]
        headers = [header for _, header in columns]
        rows = []
        for row in values[header_index + 1:]:
            record = {header: str(row[index]).strip() for index, header in columns if index < len(row) and row[index] not in (None, '')}
            if record:
                rows.append(record)
        workbook.close()
        needle = query.strip().lower()
        if needle:
            rows = [row for row in rows if needle in ' '.join(row.values()).lower()]
        return {'files': [{'name': path.name, 'status': '불러옴'}], 'locked': False, 'fields': headers, 'rows': rows, 'message': f'{path.name} · {len(rows)}건'}
    except Exception as error:
        return {'files': [{'name': path.name, 'status': '불러오기 실패'}], 'locked': False, 'fields': [], 'rows': [], 'message': f'조회명단을 읽지 못했습니다: {error}'}


def cell_at(table, col, row):
    for cell in table.findall('.//hp:tc', NS):
        address = cell.find('hp:cellAddr', NS)
        if address is not None and address.get('colAddr') == str(col) and address.get('rowAddr') == str(row):
            return cell
    raise ValueError(f'Missing original-form cell ({col}, {row})')


def put(table, col, row, value):
    if value in (None, ''):
        return
    cell = cell_at(table, col, row)
    text = cell.find('.//hp:t', NS)
    created = text is None
    if text is None:
        # Empty value cells in the supplied form contain a run but no text node.
        # Add text to that existing run so the original table layout is preserved.
        run = cell.find('.//hp:run', NS)
        if run is None:
            raise ValueError(f'Missing text run at ({col}, {row})')
        text = ET.SubElement(run, f"{{{NS['hp']}}}t")
    # A blank value cell must start on its own first line.  Adding a leading
    # newline put values (notably 신청부수) beneath the neighbouring label cell.
    text.text = str(value) if created else (text.text or '') + '\n' + str(value)


def put_centered(table, col, row, value):
    """Keep the printed label and place a short value in the centre of its cell."""
    if value in (None, ''):
        return
    cell = cell_at(table, col, row)
    text = cell.find('.//hp:t', NS)
    if text is None:
        put(table, col, row, value)
        return
    text.text = (text.text or '') + '\n' + (' ' * 12) + str(value)


def put_area(table, col, row, value):
    """Keep the entered area immediately beside the template's single ㎡ mark."""
    if value in (None, ''):
        return
    cell = cell_at(table, col, row)
    texts = cell.findall('.//hp:t', NS)
    if not texts:
        raise ValueError('Missing area text cell')
    # The supplied HWPX keeps "연면적" and "㎡" in separate text runs.
    # Put the value in the unit run so it reads naturally as "25026.07 ㎡".
    unit = next((text for text in texts if '㎡' in (text.text or '')), None)
    if unit is None:
        unit = texts[-1]
    # 별지 제12호서식 is a broad blank cell: preserve indentation so the value
    # is centred in the area field rather than placed below its label.
    unit.text = f"{' ' * 23 if (col, row) == (2, 8) else ''}{value} ㎡"


def checked(table, col, row):
    text = cell_at(table, col, row).find('.//hp:t', NS)
    text.text = (text.text or '').replace('[  ]', '[√]').replace('[ ]', '[√]')


def date_kr(value):
    try:
        year, month, day = value.split('-')
        return f'{year}년 {int(month)}월 {int(day)}일'
    except Exception:
        return value or ''


def put_date_line(table, col, row, value):
    """The original form has one cell containing year/month/day placeholders."""
    if not value:
        return
    try:
        year, month, day = value.split('-')
        # The form prints 년·월·일 as labels.  Place each number immediately
        # before its label, not underneath it.
        display = f'{year}      년     {int(month)}      월     {int(day)}      일'
    except ValueError:
        display = value
    text = cell_at(table, col, row).find('.//hp:t', NS)
    if text is None:
        raise ValueError(f'Missing date text at ({col}, {row})')
    text.text = display


def fill(template, data, fields, appointment=False):
    buffer = BytesIO()
    with zipfile.ZipFile(template, 'r') as source, zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as destination:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == 'Contents/section0.xml':
                root = ET.fromstring(content)
                table = root.findall('.//hp:tbl', NS)[0]
                for col, row, name in fields:
                    if name == 'reportDate':
                        put_date_line(table, col, row, data.get(name, ''))
                        continue
                    value = date_kr(data.get(name, '')) if name in {'managerBirth', 'appointmentDate', 'reportDate'} else data.get(name, '')
                    if name == 'buildingArea':
                        put_area(table, col, row, value)
                        continue
                    if name == 'buildingUse' and (col, row) == (8, 8):
                        put_centered(table, col, row, value)
                        continue
                    put(table, col, row, value)
                if appointment:
                    mode = data.get('appointmentType', '')
                    if mode == '직접 선임': checked(table, 8, 8)
                    elif mode == '위탁 수행': checked(table, 8, 9)
                content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            destination.writestr(info, content)
    return buffer.getvalue()


def fill_form10(data):
    # Supplied 별지 제10호서식, table 0. Different address/name cells stay distinct.
    required = ('ownerName', 'ownerRepresentative', 'businessNumber', 'ownerAddress', 'ownerPhone',
                'buildingArea', 'buildingUse', 'buildingAddress', 'managerName', 'managerBirth',
                'managerAddress', 'managerGrade', 'appointmentDate', 'licenseNumber', 'reportDate',
                'appointmentType')
    missing = [name for name in required if not str(data.get(name, '')).strip()]
    if missing:
        raise ValueError('필수 입력값이 비어 있어 빈 서식은 생성하지 않습니다: ' + ', '.join(missing))
    fields = [
        (8, 4, 'ownerName'), (20, 4, 'ownerRepresentative'), (34, 4, 'businessNumber'),
        (8, 5, 'ownerAddress'), (34, 5, 'ownerPhone'),
        (12, 6, 'buildingArea'), (30, 6, 'buildingUse'), (8, 7, 'buildingAddress'),
        (8, 11, 'managerName'), (19, 11, 'managerBirth'), (8, 12, 'managerAddress'),
        (8, 13, 'managerGrade'), (19, 13, 'appointmentDate'), (33, 13, 'licenseNumber'),
        (0, 18, 'reportDate'),
        # The signature line at the bottom is a separate blank cell.
        (25, 19, 'ownerName'),
    ]
    return fill(template_for('form10'), data, fields, appointment=True)


def fill_form12(data):
    # Supplied 별지 제12호서식, table 0.
    fields = [
        (1, 4, 'ownerName'), (8, 4, 'ownerRepresentative'), (1, 5, 'businessNumber'),
        (8, 5, 'ownerPhone'), (1, 6, 'ownerAddress'),
        (2, 8, 'buildingArea'), (8, 8, 'buildingUse'), (1, 9, 'buildingAddress'),
        # Keep the printed labels ("신청사유", "신청부수") intact and write only
        # into their adjacent blank value cells.
        (1, 11, 'certificateReason'), (7, 11, 'certificateCopies'),
        (0, 13, 'reportDate'), (0, 14, 'ownerName'),
    ]
    return fill(template_for('form12'), data, fields)


FONT_PATH = Path(r'C:\Windows\Fonts\malgun.ttf')


def preview_font(size):
    """Use a Korean-capable font when producing an in-browser form preview."""
    try:
        return ImageFont.truetype(str(FONT_PATH), size)
    except OSError:
        return ImageFont.load_default()


def draw_value(draw, position, value, max_width, size=10):
    """Write a generated field into the matching blank cell of the supplied form."""
    value = str(value or '')
    font_size = size
    font = preview_font(font_size)
    while font_size > 7 and draw.textbbox((0, 0), value, font=font)[2] > max_width:
        font_size -= 1
        font = preview_font(font_size)
    draw.text(position, value, fill=(0, 0, 0), font=font)


def draw_date_parts(draw, value, positions):
    """The printed form already says 년/월/일, so write only the three numbers."""
    try:
        year, month, day = value.split('-')
        for position, part in zip(positions, (year, str(int(month)), str(int(day)))):
            draw_value(draw, position, part, 38)
    except ValueError:
        draw_value(draw, positions[0], value, 130)


def generated_preview(kind, data):
    """Render the same input that will be inserted into the downloadable HWPX.

    HWPX preview PNGs embedded in the original templates are intentionally blank.
    This overlays the generated values on those exact supplied form images, so the
    browser preview is a faithful view of the file about to be downloaded.
    """
    with zipfile.ZipFile(template_for(kind)) as source:
        image = Image.open(BytesIO(source.read('Preview/PrvImage.png'))).convert('RGB')
    draw = ImageDraw.Draw(image)
    if kind == 'form10':
        fields = [
            ((190, 183), 'ownerName', 165), ((360, 183), 'ownerRepresentative', 150), ((518, 183), 'businessNumber', 130),
            ((190, 216), 'ownerAddress', 320), ((518, 216), 'ownerPhone', 130),
            ((370, 252), 'buildingArea', 54), ((465, 252), 'buildingUse', 180), ((190, 304), 'buildingAddress', 455),
            ((190, 421), 'managerName', 150), ((348, 421), 'managerBirth', 300), ((190, 454), 'managerAddress', 455),
            ((190, 488), 'managerGrade', 150), ((348, 488), 'appointmentDate', 150), ((508, 488), 'licenseNumber', 135),
            ((445, 679), 'ownerName', 135),
        ]
        for position, name, width in fields:
            value = date_kr(data.get(name, '')) if name in {'managerBirth', 'appointmentDate'} else data.get(name, '')
            draw_value(draw, position, value, width)
        if data.get('appointmentType') == '직접 선임':
            draw.text((204, 342), '✓', fill=(0, 0, 0), font=preview_font(12))
        elif data.get('appointmentType') == '위탁 수행':
            draw.text((204, 371), '✓', fill=(0, 0, 0), font=preview_font(12))
        # The template already prints 년·월·일.  Keep the numbers on that same
        # baseline, immediately before each label, just as in the HWPX form.
        draw_date_parts(draw, data.get('reportDate', ''), ((486, 646), (550, 646), (620, 646)))
    else:
        fields = [
            ((150, 188), 'ownerName', 300), ((462, 188), 'ownerRepresentative', 180),
            ((150, 240), 'businessNumber', 300), ((462, 240), 'ownerPhone', 180), ((150, 293), 'ownerAddress', 485),
            # Align the value immediately to the left of the supplied ㎡ mark.
            ((350, 368), 'buildingArea', 48), ((515, 375), 'buildingUse', 100), ((150, 433), 'buildingAddress', 485),
            ((150, 546), 'certificateReason', 235), ((475, 530), 'certificateCopies', 120), ((442, 680), 'ownerName', 135),
        ]
        for position, name, width in fields:
            value = data.get(name, '')
            draw_value(draw, position, value, width)
        # Numbers sit on the same line, immediately to the left of 년·월·일.
        draw_date_parts(draw, data.get('reportDate', ''), ((510, 632), (580, 632), (622, 632)))
    output = BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


class Handler(BaseHTTPRequestHandler):
    def cors_origin(self):
        """Allow the two local addresses used when opening the portal."""
        origin = self.headers.get('Origin', '')
        if origin in ('http://127.0.0.1:8080', 'http://localhost:8080'):
            return origin
        return 'http://127.0.0.1:8080'

    def send_response_headers(self, length=0, name=''):
        self.send_header('Access-Control-Allow-Origin', self.cors_origin())
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        if length:
            self.send_header('Content-Type', 'application/vnd.hancom.hwpx')
            self.send_header('Content-Length', str(length))
            self.send_header('Content-Disposition', f'attachment; filename="{name}"')

    def do_OPTIONS(self):
        self.send_response(204); self.send_response_headers(); self.end_headers()

    def do_GET(self):
        if self.path.startswith('/admin/duplicate-registry'):
            query=parse_qs(urlparse(self.path).query).get('q', [''])[0]
            payload=json.dumps(duplicate_registry(query),ensure_ascii=False).encode('utf-8')
            self.send_response(200); self.send_header('Access-Control-Allow-Origin', self.cors_origin()); self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Content-Length',str(len(payload))); self.end_headers(); self.wfile.write(payload); return
        if self.path.startswith('/admin/template/'):
            kind = self.path.rsplit('/', 1)[-1]
            if kind not in ('draft', 'certificate'):
                self.send_error(404); return
            content = admin_document(kind)
            # HTTP header values are latin-1 in the standard library.  Keep the
            # downloaded filename ASCII while the page itself presents Korean
            # document names.
            name = 'draft-preapproval.hwpx' if kind == 'draft' else 'certificate-application.hwpx'
            self.send_response(200); self.send_response_headers(len(content), name); self.end_headers(); self.wfile.write(content); return
        if self.path.startswith('/admin/preview/'):
            kind = self.path.rsplit('/', 1)[-1]
            if kind not in ('draft', 'certificate'):
                self.send_error(404); return
            content = admin_preview(kind)
            self.send_response(200); self.send_header('Access-Control-Allow-Origin', self.cors_origin()); self.send_header('Content-Type', 'image/png'); self.send_header('Content-Length', str(len(content))); self.end_headers(); self.wfile.write(content); return
        previews = {'/preview/form10': 'form10', '/preview/form12': 'form12'}
        if self.path not in previews:
            self.send_error(404); return
        with zipfile.ZipFile(template_for(previews[self.path])) as source:
            content = source.read('Preview/PrvImage.png')
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', self.cors_origin())
        self.send_header('Content-Type', 'image/png')
        self.send_header('Content-Length', str(len(content)))
        self.end_headers(); self.wfile.write(content)

    def do_POST(self):
        try:
            data = json.loads(self.rfile.read(int(self.headers.get('Content-Length', '0'))))
            # A report is dated on the day it is created, not on the appointment date.
            if self.path.startswith(('/generate/form', '/preview/generated/form')):
                data['reportDate'] = datetime.now(ZoneInfo('Asia/Seoul')).date().isoformat()
            if self.path == '/qa/ask':
                question = str(data.get('question', '')).strip()
                if not question:
                    payload = json.dumps({'error': '질문을 입력해 주세요.'}, ensure_ascii=False).encode('utf-8')
                    self.send_response(400); self.send_header('Access-Control-Allow-Origin', self.cors_origin()); self.send_header('Content-Type', 'application/json; charset=utf-8'); self.send_header('Content-Length', str(len(payload))); self.end_headers(); self.wfile.write(payload); return
                answer, sources = ask_qa(question)
                payload = json.dumps({'answer': answer, 'sources': sources}, ensure_ascii=False).encode('utf-8')
                self.send_response(200); self.send_header('Access-Control-Allow-Origin', self.cors_origin()); self.send_header('Content-Type', 'application/json; charset=utf-8'); self.send_header('Content-Length', str(len(payload))); self.end_headers(); self.wfile.write(payload); return
            if self.path == '/generate/form10':
                content, name = fill_form10(data), 'form10-filled.hwpx'
            elif self.path == '/generate/form12':
                content, name = fill_form12(data), 'form12-filled.hwpx'
            elif self.path.startswith('/admin/document/'):
                kind = self.path.rsplit('/', 1)[-1]
                if kind not in ('draft', 'certificate'):
                    self.send_error(404); return
                content = admin_document(kind, data)
                name = 'draft-preapproval.hwpx' if kind == 'draft' else 'certificate-application.hwpx'
            elif self.path.startswith('/admin/preview/'):
                kind = self.path.rsplit('/', 1)[-1]
                if kind not in ('draft', 'certificate'):
                    self.send_error(404); return
                content = admin_preview(kind, data)
                self.send_response(200); self.send_header('Access-Control-Allow-Origin', self.cors_origin()); self.send_header('Content-Type', 'image/png'); self.send_header('Content-Length', str(len(content))); self.end_headers(); self.wfile.write(content); return
            elif self.path == '/preview/generated/form10':
                fill_form10(data)  # preview only a document that can really be generated
                content = generated_preview('form10', data)
                self.send_response(200); self.send_header('Access-Control-Allow-Origin', self.cors_origin()); self.send_header('Content-Type', 'image/png'); self.send_header('Content-Length', str(len(content))); self.end_headers(); self.wfile.write(content); return
            elif self.path == '/preview/generated/form12':
                content = generated_preview('form12', data)
                self.send_response(200); self.send_header('Access-Control-Allow-Origin', self.cors_origin()); self.send_header('Content-Type', 'image/png'); self.send_header('Content-Length', str(len(content))); self.end_headers(); self.wfile.write(content); return
            else:
                self.send_error(404); return
            self.send_response(200); self.send_response_headers(len(content), name); self.end_headers(); self.wfile.write(content)
        except QAConfigurationError as error:
            payload = json.dumps({'error': str(error)}, ensure_ascii=False).encode('utf-8')
            self.send_response(503); self.send_header('Access-Control-Allow-Origin', self.cors_origin()); self.send_header('Content-Type', 'application/json; charset=utf-8'); self.send_header('Content-Length', str(len(payload))); self.end_headers(); self.wfile.write(payload)
        except Exception as error:
            self.send_response(500); self.send_response_headers(); self.end_headers(); self.wfile.write(str(error).encode())

    def log_message(self, *_):
        pass


if __name__ == '__main__':
    HTTPServer(('127.0.0.1', 8091), Handler).serve_forever()
